import os
import PyPDF2
from docx import Document
import subprocess
import re
from openpyxl import Workbook

FOLDER_PATH = "uploaded_files"

# Extract email and phone numbers using regex
EMAIL_PATTERN = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}(?:\d+)?\b'

PHONE_PATTERNS = r'(\d{3}[-\.\s]??\d{3}[-\.\s]??\d{4}|\(\d{3}\)\s*\d{3}[-\.\s]??\d{4}|\d{3}[-\.\s]??\d{4})'


def list_files(path):
    return os.listdir(path) if os.path.isdir(path) else []

def read_pdf(file_path):
    try:
        with open(file=f"{FOLDER_PATH}/{file_path}", mode="rb") as file:
            pdf = PyPDF2.PdfReader(file)
            text = ""
            for page_num in range(len(pdf.pages)):
                page = pdf.pages[page_num]
                text += page.extract_text()
    except:
        return None        
    return text

def read_docx(file_path):
    try:
        with open(file=f"{FOLDER_PATH}/{file_path}", mode="rb") as file:
            docx = Document(file)
            text = ""
            for paragraph in docx.paragraphs:
                text += paragraph.text
    except:
        return None
    return text

def read_doc(file_path):
    try:
        process = subprocess.Popen(["antiword", f"{FOLDER_PATH}/{file_path}"], stdout=subprocess.PIPE)    
        stdout, _ = process.communicate()
        text = stdout.decode("utf-8")
        return text
    except:
        return None

def save_to_xls(data, output_file):
    wb = Workbook()
    ws = wb.active
    for row in data:
        ws.append(row)
        
    wb.save(output_file)
    
    
def main():
    data_list = [("Email", "Phone Number", "Text")]
    files = list_files(FOLDER_PATH)
    for file in files:
        extension = file[-3:]
        if extension == "pdf":            
            text = read_pdf(file)
        elif(extension == "doc"):
            text = read_doc(file)
        else:
            text = read_docx(file)
            
        email_match = re.search(EMAIL_PATTERN, text)
        email = email_match.group() if email_match else ''
        if email != "":
            if email[-1].isdigit():
                email = email[:-1]
        
        phone_match = re.search(PHONE_PATTERNS, text)
        phone_number = phone_match.group() if phone_match else ''
        
        data_list.append((email, phone_number, text))
        
    file_name = "output.xls"   
    save_to_xls(data=data_list, output_file=file_name)
              
# if __name__ == "__main__":
#     main()